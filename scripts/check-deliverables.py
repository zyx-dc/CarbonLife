# -*- coding: utf-8 -*-
"""Assert the delivery folder is submission-ready (work order section 6).

Usage:
  python scripts/check-deliverables.py [--dir <folder with the 4 deliverables>]

Note vs. the ticket's snippet: `len(d.tables) == 7` no longer holds because the
side-by-side screenshot pairs are laid out with 4 borderless layout tables. The check
below counts CONTENT tables (tables without images) = 7 and keeps images = 15, which is
what "7 张表 + 15 张图" actually means.
"""
import argparse
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from docx import Document
    from pypdf import PdfReader
except ImportError as e:
    sys.exit(f"missing dependency: {e}")

DEFAULT = Path(r"C:\Users\39748\Documents\默认项目\CarbonLife-复核-2026-10-21\3-交付件")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(DEFAULT))
    a = ap.parse_args()
    B = Path(a.dir)

    ppt = B / "CarbonLife-答辩PPT.pptx"
    doc = B / "CarbonLife-设计文档.docx"
    pdf = B / "CarbonLife-设计文档.pdf"
    for p in (ppt, doc, pdf):
        assert p.exists(), f"missing {p}"

    # ---------- PPT ----------
    pr = Presentation(str(ppt))
    assert len(pr.slides) == 15, f"PPT 页数 {len(pr.slides)}"
    texts = lambda i: " ".join(sh.text_frame.text for sh in pr.slides[i].shapes if sh.has_text_frame)
    s5 = texts(4)
    assert "10 大功能" in s5, "P5 标题不是 10 大功能"
    nums = {f"{i:02d}" for i in range(1, 11)}
    cards = [sh.text_frame.text.strip() for sh in pr.slides[4].shapes
             if sh.has_text_frame and sh.text_frame.text.strip() in nums]
    assert len(cards) == 10, f"P5 卡片数 {len(cards)}"
    s14 = texts(13)
    for kw in ("云备份端到端", "导入运行时", "跨设备家庭聚合", "分布式权限弹窗"):
        assert kw in s14, f"P14 待复测缺：{kw}"

    # ---------- DOCX ----------
    d = Document(str(doc))
    sec = d.sections[0]
    assert abs(sec.page_width.cm - 21.0) < 0.1 and abs(sec.page_height.cm - 29.7) < 0.1, \
        f"纸张不是 A4：{sec.page_width.cm:.1f}x{sec.page_height.cm:.1f}cm"
    assert len(d.inline_shapes) == 15, f"图片数 {len(d.inline_shapes)}"
    content_tables = [t for t in d.tables if "graphic" not in t._tbl.xml]
    assert len(content_tables) == 7, f"内容表数 {len(content_tables)}（总表 {len(d.tables)}）"

    t2 = content_tables[1]
    cols = [r.cells[0].text.strip() for r in t2.rows]
    assert cols == ["编号"] + [f"F{i}" for i in range(1, 11)] + ["D1"], cols

    alltext = "\n".join([p.text for p in d.paragraphs]
                        + [c.text for t in d.tables for r in t.rows for c in r.cells])
    for banned in ("F11", "实机"):
        assert banned not in alltext, f"文档仍出现「{banned}」"

    # headers repeat on the content tables (tblHeader lives inside w:trPr)
    from docx.oxml.ns import qn
    for i, t in enumerate(content_tables, start=1):
        trPr = t.rows[0]._tr.find(qn("w:trPr"))
        assert trPr is not None and trPr.find(qn("w:tblHeader")) is not None, \
            f"表 {i} 表头未设置跨页重复"

    # ---------- PDF ----------
    r = PdfReader(str(pdf))
    assert len(r.pages) >= 22, f"PDF 页数 {len(r.pages)}"
    for i, pg in enumerate(r.pages, 1):
        assert round(float(pg.mediabox.width)) == 595 and round(float(pg.mediabox.height)) == 842, \
            f"p{i} 不是 A4"
        assert len((pg.extract_text() or "").strip()) > 3, f"p{i} 是空白页"

    print(f"ALL PASS  —  {len(r.pages)} pages | 15 slides | 7 content tables + "
          f"{len(d.tables) - 7} layout tables | 15 figures")


if __name__ == "__main__":
    main()
