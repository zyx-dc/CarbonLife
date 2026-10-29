"""
把 docs/SPEECH.md 转成 docx（够用即可：标题/段落/表格/列表/引用/代码块/加粗/行内代码）。

用法：python tools/md2docx.py docs/SPEECH.md docs/CarbonLife-答辩稿.docx "CarbonLife 答辩稿"
"""
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

CJK_FONT = "微软雅黑"
CODE_FONT = "Consolas"
INK = RGBColor(0x1A, 0x1A, 0x1A)
SUB = RGBColor(0x6B, 0x6B, 0x6B)
ACCENT = RGBColor(0x1F, 0x5A, 0x3D)


def set_cjk(run, font=CJK_FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font)


def add_rich(paragraph, text, size=10.5, color=INK):
    """处理 **粗体** 与 `行内代码`。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        raw = part[2:-2] if bold else (part[1:-1] if code else part)
        run = paragraph.add_run(raw)
        run.bold = bold
        run.font.size = Pt(size)
        run.font.color.rgb = ACCENT if bold and size >= 12 else color
        set_cjk(run, CODE_FONT if code else CJK_FONT)


def convert(md_path, docx_path, title="CarbonLife 答辩稿"):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = CJK_FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)

    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = ACCENT
    set_cjk(r)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("校内答辩 · 演讲稿 + 技术框架")
    r.font.size = Pt(9)
    r.font.color.rgb = SUB
    set_cjk(r)

    lines = open(md_path, encoding="utf-8").read().splitlines()
    i = 0
    in_code = False
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if not table_buf:
            return
        rows = [r for r in table_buf if not re.match(r"^\|[\s:|-]+\|$", r.strip())]
        cells = [[c.strip() for c in row.strip().strip("|").split("|")] for row in rows]
        if not cells:
            table_buf = []
            return
        cols = max(len(r) for r in cells)
        t = doc.add_table(rows=0, cols=cols)
        t.style = "Light Grid Accent 1"
        for ri, row in enumerate(cells):
            tr = t.add_row().cells
            for ci in range(cols):
                txt = row[ci] if ci < len(row) else ""
                p = tr[ci].paragraphs[0]
                add_rich(p, txt, size=9)
                if ri == 0:
                    for run in p.runs:
                        run.bold = True
        doc.add_paragraph()
        table_buf = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_table()
            in_code = not in_code
            i += 1
            continue

        if in_code:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(line)
            run.font.size = Pt(9)
            run.font.name = CODE_FONT
            set_cjk(run, CODE_FONT)
            i += 1
            continue

        if stripped.startswith("|"):
            table_buf.append(stripped)
            i += 1
            continue
        else:
            flush_table()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped.lstrip("#").strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12 if level <= 2 else 8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt({1: 16, 2: 14, 3: 12, 4: 11}.get(level, 11))
            run.font.color.rgb = ACCENT if level <= 2 else INK
            set_cjk(run)
            i += 1
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_rich(p, stripped.lstrip("> ").strip(), size=9.5, color=SUB)
            for run in p.runs:
                run.italic = True
            i += 1
            continue

        if re.match(r"^[-*]\s+", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_rich(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        if re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_rich(p, re.sub(r"^\d+\.\s+", "", stripped))
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        add_rich(p, stripped)
        i += 1

    flush_table()
    doc.save(docx_path)
    print("saved:", docx_path)


if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2]
    title = sys.argv[3] if len(sys.argv) > 3 else "CarbonLife 答辩稿"
    convert(src, dst, title)
